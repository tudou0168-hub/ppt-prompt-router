#!/usr/bin/env python3
"""Cross-platform installer for ppt-prompt-router itself."""

from __future__ import annotations

import argparse
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import hashlib
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path
from typing import Any

from installers import get_adapter, host_choices
from installers.base import HostError


PACKAGE_NAME = "ppt-prompt-router"
REQUIRED_ROOT_FILES = (
    "SKILL.md", "README.md", "VERSION", "prompt-index.json", "install.py", "CHANGELOG.md", ".gitignore",
)
REQUIRED_ROOT_DIRS = ("prompts", "installers", "scripts", "integrations", "docs", "tests", ".github")
SUPPORTED_INDEX_SCHEMAS = {"1.0", "2.0"}
EXPECTED_CATEGORY_COUNTS = {
    "01_workplace_office": 7,
    "02_student_campus": 4,
    "03_promotion_showcase": 4,
    "04_personal_general": 4,
    "05_professional_scenarios": 6,
}
PROFILE_FIELDS = (
    "communication_job",
    "audience_register",
    "required_story_beats",
    "required_content_fields",
    "title_voice",
    "opening_task",
    "closing_task",
    "visual_posture",
    "evidence_standard",
    "key_failure_modes",
)
PROFILE_LIST_FIELDS = {"required_story_beats", "required_content_fields", "key_failure_modes"}
PROFILE_DEFAULTS: dict[str, Any] = {
    "communication_job": "帮助受众理解材料重点，并形成与场景匹配的下一步判断或行动。",
    "audience_register": "专业、清晰、面向真实决策或学习场景。",
    "required_story_beats": ["背景与问题", "核心判断", "路径或行动"],
    "required_content_fields": [],
    "title_voice": "面向受众的结论式标题。",
    "opening_task": "建立主题、受众关切和本次材料要回答的问题。",
    "closing_task": "收束核心结论并明确下一步。",
    "visual_posture": "专业、克制，以内容关系为中心。",
    "evidence_standard": "事实、数据、案例和承诺必须来自材料或明确标注边界。",
    "key_failure_modes": ["机械拆页", "重复卡片化", "把推断写成事实"],
}
FOCUSED_PROFILE_IDS = {
    "government_strategy",
    "business_bid",
    "data_report",
    "classroom_lesson",
    "brand_presentation",
    "fundraising_bp",
}


class PackageError(RuntimeError):
    pass


def expand_path(raw: str) -> Path:
    return Path(os.path.expandvars(os.path.expanduser(raw))).resolve(strict=False)


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def package_root(source_root: Path | None = None) -> Path:
    return (source_root or Path(__file__).resolve().parent).resolve(strict=False)


def destination_root(target_dir: Path) -> Path:
    return target_dir / PACKAGE_NAME


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
    except ValueError:
        return False
    return True


def source_target_conflict(source_root: Path, target_dir: Path, dest_root: Path) -> bool:
    source = source_root.resolve(strict=False)
    return any(
        source == path.resolve(strict=False)
        or _is_relative_to(source, path.resolve(strict=False))
        or _is_relative_to(path.resolve(strict=False), source)
        for path in (target_dir, dest_root)
    )


def load_index(root: Path) -> dict:
    path = root / "prompt-index.json"
    if not path.is_file():
        raise PackageError(f"missing prompt-index.json: {path}")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise PackageError(f"invalid prompt-index.json: {exc}") from exc


def prompt_body(root: Path, entry: dict) -> str:
    prompt_path = root / Path(entry["file"])
    if not prompt_path.is_file():
        raise PackageError(f"missing prompt file: {entry['file']}")
    return prompt_path.read_text(encoding="utf-8")


def _profile_value(value: str, field: str) -> Any:
    if field in PROFILE_LIST_FIELDS:
        return [item.strip() for item in value.split("|") if item.strip()]
    return value.strip()


def profile_fields(root: Path, entry: dict) -> tuple[dict[str, Any], set[str]]:
    """Read the lightweight Router Profile block and fill public defaults."""
    text = prompt_body(root, entry)
    lines = text.splitlines()
    in_block = False
    explicit: dict[str, Any] = {}
    for line in lines:
        if line.strip() == "## Router Profile":
            in_block = True
            continue
        if in_block and line.startswith("## "):
            break
        if not in_block or ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = key.strip().lstrip("-").strip()
        if key in PROFILE_FIELDS and value.strip():
            explicit[key] = _profile_value(value, key)
    fields = {key: explicit.get(key, PROFILE_DEFAULTS[key]) for key in PROFILE_FIELDS}
    return fields, set(explicit)


def lint_profiles(root: Path, entries: list[dict]) -> None:
    for entry in entries:
        fields, explicit = profile_fields(root, entry)
        if entry["id"] in FOCUSED_PROFILE_IDS and explicit != set(PROFILE_FIELDS):
            missing = sorted(set(PROFILE_FIELDS) - explicit)
            raise PackageError(f"focused profile {entry['id']} missing fields: {', '.join(missing)}")
        if len(fields["key_failure_modes"]) > 3:
            raise PackageError(f"profile {entry['id']} has more than 3 key_failure_modes")
        required_nonempty = [field for field in PROFILE_FIELDS if field != "required_content_fields"]
        if not all(fields[field] for field in required_nonempty):
            raise PackageError(f"profile {entry['id']} has an empty compiled field")


def compile_director_profile(root: Path, entry: dict) -> str:
    """Compile shared direction plus one scenario profile into the handoff file."""
    fields, explicit = profile_fields(root, entry)
    lint_profiles(root, [entry])
    protocol_path = root / "prompts" / "_shared" / "director_protocol.md"
    if not protocol_path.is_file():
        raise PackageError(f"missing shared director protocol: {protocol_path}")

    lines = [
        "# Director Profile",
        "",
        f"- profile_id: {entry['id']}",
        f"- validation_status: {entry.get('validation_status', 'experimental')}",
        "",
        "## Universal Director Protocol",
        "",
        protocol_path.read_text(encoding="utf-8").strip(),
        "",
        "## Scene Profile",
        "",
    ]
    for field in PROFILE_FIELDS:
        value = fields[field]
        rendered = " | ".join(value) if field in PROFILE_LIST_FIELDS else value
        lines.append(f"- {field}: {rendered}")
    if entry["id"] not in FOCUSED_PROFILE_IDS:
        lines.extend([
            "",
            "## Experimental Source Guidance",
            "",
            prompt_body(root, entry).strip(),
        ])
    return "\n".join(lines).rstrip() + "\n"


def prompt_entries(index: dict) -> list[dict]:
    prompts = index.get("prompts")
    if isinstance(prompts, list):
        return [dict(entry) for entry in prompts]
    if isinstance(prompts, dict):
        return [{"id": prompt_id, **value} for prompt_id, value in prompts.items() if isinstance(value, dict)]
    raise PackageError("prompt-index prompts must be a list or object")


def validate_prompt_index(index: dict) -> list[dict]:
    if str(index.get("schema_version", "1.0")) not in SUPPORTED_INDEX_SCHEMAS:
        raise PackageError("prompt-index schema_version must be 1.0 or 2.0")
    entries = prompt_entries(index)
    if len(entries) != 25:
        raise PackageError("prompt-index must contain exactly 25 prompts")
    ids, files = set(), set()
    categories = {key: 0 for key in EXPECTED_CATEGORY_COUNTS}
    for entry in entries:
        for key in ("id", "name_zh", "category", "file", "use_when", "strong_signals", "avoid_when", "conflicts_with"):
            if key not in entry or not entry[key]:
                raise PackageError(f"prompt-index entry missing {key}")
        if entry["id"] in ids or entry["file"] in files:
            raise PackageError("prompt-index contains duplicate id or file")
        if entry["category"] not in categories:
            raise PackageError(f"unexpected category: {entry['category']}")
        ids.add(entry["id"])
        files.add(entry["file"])
        categories[entry["category"]] += 1
    if categories != EXPECTED_CATEGORY_COUNTS:
        raise PackageError(f"category counts mismatch: {categories}")
    return entries


def select_prompt(index: dict, task: str, *, prompt_id: str | None = None) -> dict | None:
    """Return one matching prompt, or None when the task is genuinely unclear."""
    entries = validate_prompt_index(index)
    if prompt_id:
        return next((entry for entry in entries if entry["id"] == prompt_id), None)
    text = task.lower()
    ranked: list[tuple[int, str, dict]] = []
    for entry in entries:
        terms = [entry["id"], entry["name_zh"], *entry.get("required_signals", []), *entry.get("strong_signals", []), *entry.get("supporting_signals", [])]
        score = sum(1 for term in terms if str(term).lower() in text)
        if score:
            ranked.append((score, entry["id"], entry))
    if not ranked:
        return None
    ranked.sort(key=lambda item: (-item[0], item[1]))
    return ranked[0][2]


def copy_package_tree(source_root: Path, staging_root: Path) -> None:
    entries = validate_prompt_index(load_index(source_root))
    staging_root.mkdir(parents=True, exist_ok=False)
    for name in REQUIRED_ROOT_FILES:
        source = source_root / name
        if not source.is_file():
            raise PackageError(f"missing required file: {name}")
        shutil.copy2(source, staging_root / name)

    for entry in entries:
        relative = Path(entry["file"])
        source = source_root / relative
        destination = staging_root / relative
        if not source.is_file():
            raise PackageError(f"missing prompt file: {relative}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
    shared_protocol = source_root / "prompts" / "_shared" / "director_protocol.md"
    if not shared_protocol.is_file():
        raise PackageError("missing prompts/_shared/director_protocol.md")
    shared_destination = staging_root / "prompts" / "_shared"
    shared_destination.mkdir(parents=True, exist_ok=True)
    shutil.copy2(shared_protocol, shared_destination / shared_protocol.name)

    adapters = source_root / "installers"
    if not adapters.is_dir():
        raise PackageError("missing host adapters directory")
    shutil.copytree(
        adapters,
        staging_root / "installers",
        ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "*.pyo", ".DS_Store", ".git"),
    )
    for name in ("scripts", "integrations", "docs", "tests", ".github"):
        source = source_root / name
        if not source.is_dir():
            raise PackageError(f"missing package directory: {name}")
        shutil.copytree(
            source,
            staging_root / name,
            ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "*.pyo", ".DS_Store", ".git", "projects", "*.pptx", "*.docx", "*.pdf"),
        )


def validate_package_root(root: Path) -> None:
    root = root.resolve(strict=False)
    if not root.is_dir():
        raise PackageError(f"package root not found: {root}")
    expected_top = set(REQUIRED_ROOT_FILES) | set(REQUIRED_ROOT_DIRS)
    actual_top = {path.name for path in root.iterdir()}
    allowed_local_extras = {"projects", ".DS_Store", ".codex", "__pycache__", "COMMERCIAL_LICENSE.md", "THIRD_PARTY_NOTICES.md"}
    unexpected = actual_top - expected_top - allowed_local_extras
    if expected_top - actual_top or unexpected:
        raise PackageError(f"top-level structure mismatch; missing={sorted(expected_top - actual_top)}, extra={sorted(unexpected)}")
    for name in REQUIRED_ROOT_FILES:
        path = root / name
        if not path.is_file() or not path.read_text(encoding="utf-8").strip():
            raise PackageError(f"missing or empty root file: {name}")
    version = (root / "VERSION").read_text(encoding="utf-8").strip()
    if f"version: {version}" not in (root / "SKILL.md").read_text(encoding="utf-8"):
        raise PackageError("VERSION and SKILL.md frontmatter version mismatch")

    entries = validate_prompt_index(load_index(root))
    lint_profiles(root, entries)
    expected_paths = {entry["file"] for entry in entries}
    actual_paths = {
        str(path.relative_to(root).as_posix())
        for path in (root / "prompts").rglob("*.md")
        if "_shared" not in path.relative_to(root / "prompts").parts
    }
    if actual_paths != expected_paths:
        raise PackageError("prompt tree does not match prompt-index.json")
    for relative in expected_paths:
        if not (root / relative).read_text(encoding="utf-8").strip():
            raise PackageError(f"empty prompt file: {relative}")
    if not (root / "prompts" / "_shared" / "director_protocol.md").is_file():
        raise PackageError("missing shared director protocol")
    for name in ("__init__.py", "base.py", "claude_code.py", "codex.py", "openclaw.py", "hermes.py", "generic.py"):
        if not (root / "installers" / name).is_file():
            raise PackageError(f"missing host adapter: installers/{name}")
    if not (root / "scripts" / "route.py").is_file():
        raise PackageError("missing Router 2.1 entrypoint: scripts/route.py")
    manifest = root / "integrations" / "ppt-master" / "manifest.json"
    lock = root / "integrations" / "ppt-master" / "upstream.lock"
    if not manifest.is_file() or not lock.is_file():
        raise PackageError("missing PPT Master integration manifest or upstream lock")
    try:
        payload = json.loads(manifest.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise PackageError(f"invalid PPT Master manifest: {exc}") from exc
    if not isinstance(payload.get("files"), list) or not payload["files"]:
        raise PackageError("PPT Master manifest must contain overlay files")
    for item in payload["files"]:
        relative = item.get("path")
        if not isinstance(relative, str) or not relative.startswith("skills/ppt-master/"):
            raise PackageError("invalid PPT Master overlay path")
        overlay = root / "integrations" / "ppt-master" / "overlay" / relative
        if not overlay.is_file():
            raise PackageError(f"missing overlay file: {relative}")
        if item.get("overlay_hash") != sha256_file(overlay):
            raise PackageError(f"overlay hash mismatch: {relative}")


def find_ppt_master_root(explicit: str | None = None) -> Path:
    candidates: list[Path] = []
    if explicit:
        candidates.append(expand_path(explicit))
    if os.environ.get("PPT_MASTER_ROOT"):
        candidates.append(expand_path(os.environ["PPT_MASTER_ROOT"]))
    candidates.extend([
        Path.home() / "Documents" / "ppt-master",
        package_root().parent / "ppt-master",
        Path.cwd() / "ppt-master",
    ])
    seen: set[Path] = set()
    for candidate in candidates:
        candidate = candidate.resolve(strict=False)
        if candidate in seen:
            continue
        seen.add(candidate)
        script = candidate / "skills" / "ppt-master" / "scripts" / "project_manager.py"
        if script.is_file():
            return candidate
    raise PackageError("ppt-master not found; pass --ppt-master-root or set PPT_MASTER_ROOT")


def find_ppt_master_skill_dir(explicit: str | None = None) -> Path:
    """Locate either a full PPT Master worktree or a host's installed skill."""
    candidates: list[Path] = []
    if explicit:
        candidates.append(expand_path(explicit))
    if os.environ.get("PPT_MASTER_ROOT"):
        candidates.append(expand_path(os.environ["PPT_MASTER_ROOT"]))
    candidates.extend([
        Path.home() / "Documents" / "ppt-master",
        package_root().parent / "ppt-master",
        Path.cwd() / "ppt-master",
        Path.home() / ".claude" / "skills",
        Path.home() / ".agents" / "skills",
        Path.home() / ".hermes" / "skills",
    ])
    for candidate in dict.fromkeys(path.resolve(strict=False) for path in candidates):
        for skill_dir in (
            candidate / "skills" / "ppt-master",
            candidate / "ppt-master",
            candidate,
        ):
            if (skill_dir / "scripts" / "project_manager.py").is_file():
                return skill_dir
    raise PackageError("ppt-master skill not found; pass --ppt-master-root or set PPT_MASTER_ROOT")


def run_tool(args: list[str], *, cwd: Path | None = None) -> str:
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    env["PYTHONIOENCODING"] = "utf-8:replace"
    try:
        result = subprocess.run(
            args,
            cwd=str(cwd) if cwd else None,
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=env,
        )
    except FileNotFoundError as exc:
        raise PackageError(f"missing executable: {args[0]}") from exc
    except subprocess.CalledProcessError as exc:
        details = (exc.stderr or exc.stdout or "").strip()
        raise PackageError(details or f"command failed: {' '.join(args)}") from exc
    output = (result.stdout or "").strip()
    if output:
        print(output)
    err = (result.stderr or "").strip()
    if err:
        print(err)
    return output


def build_handoff_markdown(
    *,
    project_name: str,
    request_text: str,
    entry: dict,
    prompt_text: str,
    source_paths: list[str],
    page_count: str | None,
    audience: str | None,
    template_paths: list[str],
) -> str:
    source_lines = [f"- {path}" for path in source_paths] if source_paths else ["- none provided"]
    template_lines = [f"- {path}" for path in template_paths] if template_paths else ["- none provided"]
    lines = [
        "# Router Handoff",
        "",
        "## Project",
        "",
        f"- project_name: {project_name}",
        "",
        "## Request",
        "",
        request_text.strip(),
        "",
        "## Sources",
        "",
    ]
    lines.extend(source_lines)
    lines.extend([
        "",
        "## Templates",
        "",
    ])
    lines.extend(template_lines)
    lines.extend([
        "",
        "## Constraints",
        "",
        f"- page_count: {page_count or 'unspecified'}",
        f"- audience: {audience or 'unspecified'}",
        f"- templates: {', '.join(template_paths) if template_paths else 'unspecified'}",
        "",
        "## Selected Prompt",
        "",
        f"- prompt_id: {entry['id']}",
        f"- prompt_name: {entry['name_zh']}",
        f"- prompt_file: {entry['file']}",
        "",
        "```text",
        prompt_text.rstrip(),
        "```",
        "",
    ])
    return "\n".join(lines)


def cmd_route(args: argparse.Namespace) -> int:
    route_script = package_root() / "scripts" / "route.py"
    if not route_script.is_file():
        raise PackageError(f"Router 2.1 entrypoint not found: {route_script}")
    forwarded = [sys.executable, str(route_script)]
    for key in (
        "request", "page_count", "audience", "purpose", "prompt_id", "format",
        "project_name", "project_base", "ppt_master_root", "template_intent",
        "director_plan", "host",
    ):
        value = getattr(args, key, None)
        if value is not None:
            forwarded.extend([f"--{key.replace('_', '-')}", str(value)])
    for source in args.source:
        forwarded.extend(["--source", source])
    for template in args.template:
        forwarded.extend(["--template", template])
    if args.move:
        forwarded.append("--move")
    return subprocess.run(forwarded, check=False).returncode


def install_package(
    target_dir: Path,
    *,
    force: bool = False,
    source_root: Path | None = None,
    mode: str = "copy",
) -> Path:
    source_root = package_root(source_root)
    target_dir = target_dir.resolve(strict=False)
    destination = destination_root(target_dir)
    if source_target_conflict(source_root, target_dir, destination):
        raise PackageError("source and target paths overlap; refusing recursive install")
    if destination.exists() and not force:
        raise PackageError(f"destination already exists: {destination}")
    target_dir.mkdir(parents=True, exist_ok=True)
    staging_parent = Path(tempfile.mkdtemp(prefix=f".{PACKAGE_NAME}-stage-", dir=str(target_dir.parent)))
    staging = staging_parent / PACKAGE_NAME
    backup = None
    try:
        if mode == "copy":
            copy_package_tree(source_root, staging)
            validate_package_root(staging)
        elif mode == "symlink":
            validate_package_root(source_root)
            staging.symlink_to(source_root, target_is_directory=True)
        else:
            raise PackageError("mode must be copy or symlink")
        if destination.exists():
            backup = staging_parent / f"{PACKAGE_NAME}.backup"
            os.replace(destination, backup)
        os.replace(staging, destination)
        return destination
    except Exception:
        if backup and backup.exists() and not destination.exists():
            os.replace(backup, destination)
        raise
    finally:
        shutil.rmtree(staging_parent, ignore_errors=True)


def uninstall_package(target_dir: Path) -> Path:
    destination = destination_root(target_dir.resolve(strict=False))
    if destination.exists():
        shutil.rmtree(destination)
    return destination


def _integration_files(root: Path) -> list[dict[str, Any]]:
    manifest = root / "integrations" / "ppt-master" / "manifest.json"
    try:
        return list(json.loads(manifest.read_text(encoding="utf-8"))["files"])
    except (OSError, KeyError, TypeError, json.JSONDecodeError) as exc:
        raise PackageError(f"invalid PPT Master manifest: {exc}") from exc


def _upstream_lock(root: Path) -> dict[str, Any]:
    path = root / "integrations" / "ppt-master" / "upstream.lock"
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise PackageError(f"invalid PPT Master upstream lock: {exc}") from exc
    for key in ("repository", "commit", "required_capabilities"):
        if not payload.get(key):
            raise PackageError(f"PPT Master upstream lock missing {key}")
    return payload


def _relative_skill_path(item: dict[str, Any]) -> Path:
    relative = Path(item["path"])
    prefix = Path("skills") / "ppt-master"
    try:
        return relative.relative_to(prefix)
    except ValueError as exc:
        raise PackageError(f"invalid PPT Master overlay path: {relative}") from exc


def verify_master_baseline(root: Path, master_root: Path) -> dict[str, Any]:
    """Verify the exact files the overlay will replace, without requiring Git."""
    lock = _upstream_lock(root)
    missing: list[str] = []
    mismatched: list[str] = []
    for item in _integration_files(root):
        target = master_root / item["path"]
        expected = item.get("upstream_hash")
        if expected is None:
            if target.exists():
                mismatched.append(item["path"])
            continue
        if not target.is_file():
            missing.append(item["path"])
        elif sha256_file(target) != expected:
            mismatched.append(item["path"])
    if missing or mismatched:
        detail = []
        if missing:
            detail.append(f"missing={', '.join(missing)}")
        if mismatched:
            detail.append(f"hash_mismatch={', '.join(mismatched)}")
        raise PackageError("PPT Master does not match upstream.lock: " + "; ".join(detail))
    return {
        "required_commit": lock["commit"],
        "verification": "overlay upstream hashes",
        "repository": lock["repository"],
    }


def _codeload_url(lock: dict[str, Any]) -> str:
    parsed = urllib.parse.urlparse(lock["repository"])
    parts = [part for part in parsed.path.split("/") if part]
    if len(parts) < 2:
        raise PackageError("upstream.lock repository is not a GitHub repository URL")
    owner, repository = parts[-2], parts[-1].removesuffix(".git")
    return f"https://codeload.github.com/{owner}/{repository}/zip/{lock['commit']}"


def _safe_extract_zip(archive: Path, destination: Path) -> None:
    try:
        with zipfile.ZipFile(archive) as package:
            for member in package.infolist():
                target = (destination / member.filename).resolve(strict=False)
                try:
                    target.relative_to(destination.resolve())
                except ValueError as exc:
                    raise PackageError("PPT Master ZIP contains an unsafe path") from exc
            package.extractall(destination)
    except zipfile.BadZipFile as exc:
        raise PackageError("PPT Master ZIP is invalid") from exc


def _find_extracted_master_root(destination: Path) -> Path:
    candidates = [destination, *sorted(path for path in destination.iterdir() if path.is_dir())]
    for candidate in candidates:
        if (candidate / "skills" / "ppt-master" / "scripts" / "project_manager.py").is_file():
            return candidate
    raise PackageError("PPT Master source does not contain skills/ppt-master")


def prepare_master_source(
    root: Path,
    staging: Path,
    *,
    source_dir: str | None = None,
    source_zip: str | None = None,
    codeload_url: str | None = None,
) -> tuple[Path, dict[str, Any]]:
    """Stage a clean source copy from local directory, local ZIP, or codeload ZIP."""
    root = package_root(root)
    selected = sum(value is not None for value in (source_dir, source_zip, codeload_url))
    if selected > 1:
        raise PackageError("choose only one PPT Master source")
    source_area = staging / "master-source"
    source_area.mkdir(parents=True, exist_ok=True)
    if source_dir:
        original = expand_path(source_dir)
        if not original.is_dir():
            raise PackageError(f"PPT Master source directory not found: {original}")
        copied = source_area / "local"
        shutil.copytree(original, copied, ignore=shutil.ignore_patterns(".git", "__pycache__", ".DS_Store"))
        master = _find_extracted_master_root(copied)
        method = "local_directory"
    elif source_zip:
        archive = expand_path(source_zip)
        if not archive.is_file():
            raise PackageError(f"PPT Master source ZIP not found: {archive}")
        _safe_extract_zip(archive, source_area)
        master = _find_extracted_master_root(source_area)
        method = "local_zip"
    else:
        lock = _upstream_lock(root)
        url = codeload_url or _codeload_url(lock)
        archive = staging / "ppt-master.zip"
        try:
            with urllib.request.urlopen(url, timeout=45) as response:
                archive.write_bytes(response.read())
        except OSError as exc:
            raise PackageError(f"unable to download PPT Master codeload ZIP: {exc}") from exc
        _safe_extract_zip(archive, source_area)
        master = _find_extracted_master_root(source_area)
        method = "codeload_zip"
    verification = verify_master_baseline(root, master)
    verification["source"] = method
    return master, verification


def _installation_record(master_root: Path) -> Path:
    return master_root / ".ppt-prompt-router-install.json"


def apply_master_overlay(root: Path, master_root: Path) -> Path:
    """Apply only manifest-listed files and preserve every original for rollback."""
    root = package_root(root)
    master_root = master_root.resolve(strict=False)
    record_path = _installation_record(master_root)
    files = _integration_files(root)
    if record_path.exists():
        current = json.loads(record_path.read_text(encoding="utf-8"))
        if current.get("package_version") == (root / "VERSION").read_text(encoding="utf-8").strip():
            for item in files:
                target = master_root / item["path"]
                if not target.is_file() or sha256_file(target) != item["overlay_hash"]:
                    raise PackageError(f"existing PPT Master overlay was changed: {item['path']}")
            return record_path
        raise PackageError("PPT Master already has a different Router overlay; uninstall it first")

    verify_master_baseline(root, master_root)

    backup = master_root / ".ppt-prompt-router-backup"
    staged: list[tuple[Path, Path | None]] = []
    try:
        for item in files:
            relative = Path(item["path"])
            target = master_root / relative
            source = root / "integrations" / "ppt-master" / "overlay" / relative
            backup_file = backup / relative if target.exists() else None
            if backup_file:
                backup_file.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(target, backup_file)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
            staged.append((target, backup_file))
        record_path.write_text(json.dumps({
            "package_version": (root / "VERSION").read_text(encoding="utf-8").strip(),
            "files": [{"path": item["path"], "overlay_hash": item["overlay_hash"]} for item in files],
        }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return record_path
    except Exception:
        for target, backup_file in reversed(staged):
            if backup_file and backup_file.exists():
                shutil.copy2(backup_file, target)
            elif target.exists():
                target.unlink()
        raise


def uninstall_master_overlay(master_root: Path) -> None:
    master_root = master_root.resolve(strict=False)
    record_path = _installation_record(master_root)
    if not record_path.is_file():
        return
    record = json.loads(record_path.read_text(encoding="utf-8"))
    backup = master_root / ".ppt-prompt-router-backup"
    for item in record.get("files", []):
        relative = Path(item["path"])
        target = master_root / relative
        if target.exists() and sha256_file(target) != item["overlay_hash"]:
            raise PackageError(f"refusing to overwrite changed PPT Master file: {relative}")
        backup_file = backup / relative
        if backup_file.exists():
            shutil.copy2(backup_file, target)
        elif target.exists():
            target.unlink()
    shutil.rmtree(backup, ignore_errors=True)
    record_path.unlink(missing_ok=True)


def _runtime_record(target_dir: Path) -> Path:
    return target_dir / ".ppt-prompt-router-runtime.json"


def _runtime_master_dir(target_dir: Path) -> Path:
    return target_dir / "ppt-master"


def deploy_runtime_master(root: Path, staged_master_root: Path, target_dir: Path, verification: dict[str, Any], *, force: bool) -> Path:
    """Deploy only the final skill into the host's formal skills directory."""
    source_skill = staged_master_root / "skills" / "ppt-master"
    if not (source_skill / "scripts" / "project_manager.py").is_file():
        raise PackageError("staged PPT Master skill is incomplete")
    target_dir = target_dir.resolve(strict=False)
    target_dir.mkdir(parents=True, exist_ok=True)
    runtime = _runtime_master_dir(target_dir)
    record = _runtime_record(target_dir)
    persistent_backup = target_dir / ".ppt-prompt-router-runtime-backup"
    if runtime.exists() or runtime.is_symlink():
        if record.is_file():
            try:
                verify_runtime_integration(root, target_dir)
                return runtime
            except PackageError:
                pass
        if not force:
            raise PackageError(f"PPT Master runtime skill already exists: {runtime}; use --force after review")
        if persistent_backup.exists():
            raise PackageError("existing PPT Master runtime backup must be restored or removed before --force")
    staging_parent = Path(tempfile.mkdtemp(prefix=".ppt-master-runtime-", dir=str(target_dir.parent)))
    staged_skill = staging_parent / "ppt-master"
    backup = persistent_backup / "ppt-master"
    record_backup = persistent_backup / "runtime-record.json"
    try:
        shutil.copytree(source_skill, staged_skill, ignore=shutil.ignore_patterns("__pycache__", ".DS_Store", "*.pyc"))
        if runtime.exists() or runtime.is_symlink():
            persistent_backup.mkdir(parents=True, exist_ok=False)
            os.replace(runtime, backup)
        if record.exists():
            os.replace(record, record_backup)
        os.replace(staged_skill, runtime)
        manifest = _integration_files(root)
        payload = {
            "schema_version": "1.0",
            "required_commit": verification["required_commit"],
            "verification": verification["verification"],
            "source": verification["source"],
            "skill_dir": "ppt-master",
            "overlay_files": [
                {"path": str(_relative_skill_path(item).as_posix()), "overlay_hash": item["overlay_hash"]}
                for item in manifest
            ],
        }
        record.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return runtime
    except Exception:
        if runtime.exists() or runtime.is_symlink():
            shutil.rmtree(runtime, ignore_errors=True)
        if backup.exists() or backup.is_symlink():
            os.replace(backup, runtime)
        if record.exists():
            record.unlink()
        if record_backup.exists():
            os.replace(record_backup, record)
        shutil.rmtree(persistent_backup, ignore_errors=True)
        raise
    finally:
        shutil.rmtree(staging_parent, ignore_errors=True)


def verify_runtime_integration(root: Path, target_dir: Path) -> dict[str, Any]:
    target_dir = target_dir.resolve(strict=False)
    runtime = _runtime_master_dir(target_dir)
    record_path = _runtime_record(target_dir)
    backup = target_dir / ".ppt-prompt-router-runtime-backup"
    if not record_path.is_file():
        raise PackageError("PPT Master runtime integration record is missing")
    try:
        record = json.loads(record_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise PackageError(f"invalid PPT Master runtime integration record: {exc}") from exc
    lock = _upstream_lock(root)
    if record.get("required_commit") != lock["commit"]:
        raise PackageError("PPT Master runtime commit does not match upstream.lock")
    if not (runtime / "SKILL.md").is_file() or not (runtime / "scripts" / "project_manager.py").is_file():
        raise PackageError("PPT Master runtime skill is incomplete")
    expected = {str(_relative_skill_path(item).as_posix()): item["overlay_hash"] for item in _integration_files(root)}
    recorded = {item.get("path"): item.get("overlay_hash") for item in record.get("overlay_files", [])}
    if recorded != expected:
        raise PackageError("PPT Master runtime overlay record does not match manifest")
    for relative, digest in expected.items():
        file_path = runtime / relative
        if not file_path.is_file() or sha256_file(file_path) != digest:
            raise PackageError(f"PPT Master runtime overlay hash mismatch: {relative}")
    return {
        "status": "verified",
        "skill_path": str(runtime),
        "upstream_commit": lock["commit"],
        "overlay_files": len(expected),
    }


def _smoke_project_path(output: str) -> Path:
    for line in output.splitlines():
        if "Project initialized:" in line:
            return Path(line.split("Project initialized:", 1)[1].strip())
    raise PackageError("PPT Master smoke init did not report a project path")


def run_master_overlay_smoke(master_root: Path, workspace: Path) -> dict[str, Any]:
    """Exercise the Router-owned Master entrypoints on a clean staged source."""
    skill = master_root / "skills" / "ppt-master"
    scripts = skill / "scripts"
    required = [
        scripts / "project_manager.py",
        scripts / "director_plan.py",
        scripts / "production.py",
        scripts / "reference_elements.py",
        scripts / "pptx_intake.py",
    ]
    if any(not path.is_file() for path in required):
        raise PackageError("PPT Master smoke prerequisites are incomplete")
    compile_result = subprocess.run(
        [sys.executable, "-m", "py_compile", *(str(path) for path in required)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if compile_result.returncode:
        raise PackageError((compile_result.stderr or compile_result.stdout).strip() or "PPT Master py_compile failed")
    import_code = (
        "import sys; "
        f"sys.path.insert(0, {str(scripts)!r}); "
        "import project_manager, reference_elements, director_plan, production"
    )
    import_result = subprocess.run(
        [sys.executable, "-c", import_code],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if import_result.returncode:
        raise PackageError((import_result.stderr or import_result.stdout).strip() or "PPT Master module import failed")
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise PackageError("reference_elements smoke requires python-pptx") from exc

    workspace.mkdir(parents=True, exist_ok=True)
    reference = workspace / "reference.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Reference Elements Smoke"
    slide.placeholders[1].text = "Visual language extraction validation"
    presentation.save(reference)

    manager = scripts / "project_manager.py"
    init_result = subprocess.run(
        [sys.executable, str(manager), "init", "overlay_smoke", "--format", "ppt169", "--dir", str(workspace)],
        cwd=str(skill),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if init_result.returncode:
        raise PackageError((init_result.stderr or init_result.stdout).strip() or "PPT Master smoke init failed")
    project = _smoke_project_path(init_result.stdout)
    template = project / "sources" / "reference.pptx"
    shutil.copy2(reference, template)
    profile = project / "analysis" / "director_profile.md"
    profile.write_text("# Smoke Director Profile\n", encoding="utf-8")
    contract = {
        "schema_version": "1.0",
        "router_version": "smoke",
        "profile": {"id": "government_strategy", "sha256": sha256_file(profile)},
        "audience": "验收",
        "purpose": "验证 reference_elements 安装完整性",
        "page_count": 4,
        "template_path": str(template),
        "template_intent": "reference_elements",
    }
    (project / "analysis" / "director_contract.json").write_text(
        json.dumps(contract, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    accept = subprocess.run(
        [sys.executable, str(manager), "router-accept", str(project)],
        cwd=str(skill),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if accept.returncode or '"accepted": true' not in accept.stdout:
        raise PackageError((accept.stderr or accept.stdout).strip() or "reference_elements router-accept smoke failed")
    if not (project / "design_spec.md").is_file() or not (project / "spec_lock.md").is_file():
        raise PackageError("reference_elements smoke did not create design files")
    plan = {
        "schema_version": "2.0",
        "content_map": {
            "core_argument": "Smoke plan validates Director Plan installation.",
            "key_findings": ["reference extraction completed"],
            "root_causes": ["overlay dependency is published"],
            "recommendations": ["continue controlled production"],
        },
        "fact_boundary": {"facts": "smoke artifacts only"},
        "storyline": {"core_narrative": "Validate planning handoff.", "logic_flow": ["validate", "plan"]},
        "pages": [
            {"page_id": "P01", "headline": "Architecture", "page_goal": "Show structure", "page_role": "architecture", "key_message": "Architecture is available.", "relationship_type": "hierarchy", "visual_anchor": "architecture", "evidence_refs": ["smoke"], "rhythm_role": "anchor"},
            {"page_id": "P02", "headline": "Flow", "page_goal": "Show flow", "page_role": "process_flow", "key_message": "Flow is controlled.", "relationship_type": "process", "visual_anchor": "flow", "evidence_refs": ["smoke"], "rhythm_role": "build"},
            {"page_id": "P03", "headline": "Comparison", "page_goal": "Show contrast", "page_role": "comparison", "key_message": "Coverage is complete.", "relationship_type": "comparison", "visual_anchor": "comparison", "evidence_refs": ["smoke"], "rhythm_role": "explain"},
            {"page_id": "P04", "headline": "Roadmap", "page_goal": "Show next step", "page_role": "roadmap", "key_message": "Proceed to production.", "relationship_type": "timeline", "visual_anchor": "timeline", "evidence_refs": ["smoke"], "rhythm_role": "close"},
        ],
        "sample_pages": [
            {"page_id": "P01", "risk_type": "card_stack_risk", "reason": "Smoke risk one"},
            {"page_id": "P02", "risk_type": "relationship_density", "reason": "Smoke risk two"},
            {"page_id": "P03", "risk_type": "visual_signature", "reason": "Smoke risk three"},
        ],
    }
    plan_path = workspace / "director_plan.json"
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    director = subprocess.run(
        [sys.executable, str(manager), "director-plan", str(project), str(plan_path)],
        cwd=str(skill),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if director.returncode or not (project / "analysis" / "director_plan.json").is_file():
        raise PackageError((director.stderr or director.stdout).strip() or "director-plan smoke failed")
    return {"status": "passed", "checks": ["py_compile", "module_import", "router_accept_reference_elements", "director_plan"]}


def uninstall_runtime_master(target_dir: Path) -> None:
    target_dir = target_dir.resolve(strict=False)
    runtime = _runtime_master_dir(target_dir)
    record_path = _runtime_record(target_dir)
    backup = target_dir / ".ppt-prompt-router-runtime-backup"
    if not record_path.is_file():
        return
    record = json.loads(record_path.read_text(encoding="utf-8"))
    if not runtime.is_dir():
        raise PackageError("PPT Master runtime skill is missing; refusing uninstall")
    for item in record.get("overlay_files", []):
        current = runtime / item["path"]
        if not current.is_file() or sha256_file(current) != item["overlay_hash"]:
            raise PackageError(f"refusing to remove changed PPT Master runtime file: {item['path']}")
    shutil.rmtree(runtime)
    record_path.unlink()
    backup_skill = backup / "ppt-master"
    backup_record = backup / "runtime-record.json"
    if backup_skill.exists() or backup_skill.is_symlink():
        os.replace(backup_skill, runtime)
    if backup_record.exists():
        os.replace(backup_record, record_path)
    shutil.rmtree(backup, ignore_errors=True)


def cmd_detect(args: argparse.Namespace) -> int:
    skills_dir = expand_path(args.skills_dir) if args.skills_dir else None
    adapter = get_adapter(args.host, skills_dir=skills_dir)
    print(json.dumps({
        "host": adapter.host_id,
        "detected": adapter.detect(),
        "skills_dirs": [str(path) for path in adapter.discover_skills_dirs()],
        "platform": platform.system().lower(),
    }, ensure_ascii=False, indent=2))
    return 0


def _target_for_install(args: argparse.Namespace) -> Path:
    if args.router_target:
        return expand_path(args.router_target)
    if args.target:
        return expand_path(args.target)
    adapter = get_adapter(args.host)
    return adapter.canonical_skills_dir()


def cmd_install(args: argparse.Namespace) -> int:
    target = _target_for_install(args)
    source = package_root()
    if args.master_root and args.master_source_dir:
        raise PackageError("--master-root and --master-source-dir cannot be used together")
    destination = destination_root(target)
    rollback_dir: Path | None = None
    rollback_copy: Path | None = None
    if destination.exists() or destination.is_symlink():
        if not args.force:
            raise PackageError(f"destination already exists: {destination}")
        rollback_dir = Path(tempfile.mkdtemp(prefix=f".{PACKAGE_NAME}-rollback-", dir=str(target.parent)))
        rollback_copy = rollback_dir / PACKAGE_NAME
        if destination.is_symlink():
            rollback_copy.symlink_to(destination.readlink(), target_is_directory=True)
        else:
            shutil.copytree(destination, rollback_copy)
    destination = install_package(target, force=args.force, source_root=source, mode=args.mode)
    stage = Path(tempfile.mkdtemp(prefix=".ppt-router-master-stage-"))
    runtime_deployed = False
    try:
        source_dir = args.master_source_dir or args.master_root
        master, verification = prepare_master_source(
            destination,
            stage,
            source_dir=source_dir,
            source_zip=args.master_source_zip,
            codeload_url=args.master_codeload_url,
        )
        apply_master_overlay(destination, master)
        smoke = run_master_overlay_smoke(master, stage / "smoke")
        runtime = deploy_runtime_master(destination, master, target, verification, force=args.force)
        runtime_deployed = True
        validate_package_root(destination)
        runtime_status = verify_runtime_integration(destination, target)
    except Exception:
        if runtime_deployed:
            uninstall_runtime_master(target)
        # Restore the Router destination when PPT Master integration fails.
        if destination.exists() or destination.is_symlink():
            uninstall_package(target)
        if rollback_copy and (rollback_copy.exists() or rollback_copy.is_symlink()):
            os.replace(rollback_copy, destination)
        raise
    finally:
        if rollback_dir:
            shutil.rmtree(rollback_dir, ignore_errors=True)
        shutil.rmtree(stage, ignore_errors=True)
    print(json.dumps({
        "router": {"status": "verified", "path": str(destination)},
        "ppt_master_overlay": {
            "status": "verified",
            "upstream_commit": verification["required_commit"],
            "source": verification["source"],
            "overlay_files": len(_integration_files(destination)),
        },
        "runtime_integration": runtime_status,
        "installation_status": "COMPLETE",
        "smoke": smoke,
    }, ensure_ascii=False, indent=2))
    return 0


def cmd_validate(args: argparse.Namespace) -> int:
    target = _target_for_install(args)
    destination = destination_root(target)
    validate_package_root(destination)
    runtime_status = verify_runtime_integration(destination, target)
    print(json.dumps({
        "router": {"status": "verified", "path": str(destination)},
        "ppt_master_overlay": {
            "status": "verified",
            "upstream_commit": _upstream_lock(destination)["commit"],
            "overlay_files": len(_integration_files(destination)),
        },
        "runtime_integration": runtime_status,
    }, ensure_ascii=False, indent=2))
    return 0


def cmd_uninstall(args: argparse.Namespace) -> int:
    if not args.yes:
        raise PackageError("uninstall requires --yes")
    target = _target_for_install(args)
    uninstall_runtime_master(target)
    print(json.dumps({
        "router": {"status": "removed", "path": str(uninstall_package(target))},
        "ppt_master_overlay": {"status": "removed"},
        "runtime_integration": {"status": "removed"},
    }, ensure_ascii=False, indent=2))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="install.py")
    sub = parser.add_subparsers(dest="command", required=True)
    install = sub.add_parser("install")
    install.add_argument("--host", choices=host_choices(), default="auto")
    install.add_argument("--target", help="兼容旧版：Agent skills parent directory")
    install.add_argument("--router-target", help="Router 安装到的技能父目录")
    install.add_argument("--master-root", help="兼容旧版：本地 PPT Master 源码目录")
    source_group = install.add_mutually_exclusive_group()
    source_group.add_argument("--master-source-dir", help="本地 PPT Master 源码目录，只读使用")
    source_group.add_argument("--master-source-zip", help="本地 PPT Master ZIP，只读使用")
    source_group.add_argument("--master-codeload-url", help="指定 codeload ZIP 地址；默认按 upstream.lock 下载")
    install.add_argument("--mode", choices=["copy", "symlink"], default="copy")
    install.add_argument("--force", action="store_true")
    install.set_defaults(func=cmd_install)
    validate = sub.add_parser("validate")
    validate.add_argument("--host", choices=host_choices(), default="auto")
    validate.add_argument("--target", help="兼容旧版：Agent skills parent directory")
    validate.add_argument("--router-target", help="Router 安装到的技能父目录")
    validate.add_argument("--master-root", help="兼容参数；运行时验证不读取此目录")
    validate.set_defaults(func=cmd_validate)
    uninstall = sub.add_parser("uninstall")
    uninstall.add_argument("--host", choices=host_choices(), default="auto")
    uninstall.add_argument("--target", help="兼容旧版：Agent skills parent directory")
    uninstall.add_argument("--router-target", help="Router 安装到的技能父目录")
    uninstall.add_argument("--master-root", help="兼容参数；卸载只恢复正式 skills 目录")
    uninstall.add_argument("--yes", action="store_true")
    uninstall.set_defaults(func=cmd_uninstall)
    detect = sub.add_parser("detect")
    detect.add_argument("--host", choices=host_choices(), default="auto")
    detect.add_argument("--skills-dir")
    detect.set_defaults(func=cmd_detect)
    route = sub.add_parser("route")
    route.add_argument("--request", required=True, help="Original user request text")
    route.add_argument("--source", action="append", default=[], help="Source file paths")
    route.add_argument("--template", action="append", default=[], help="Template file paths")
    route.add_argument("--page-count")
    route.add_argument("--audience")
    route.add_argument("--purpose")
    route.add_argument("--prompt-id")
    route.add_argument("--format", default="ppt169")
    route.add_argument("--project-name")
    route.add_argument("--project-base")
    route.add_argument("--ppt-master-root")
    route.add_argument("--template-intent", choices=["reference_elements", "native_fill", "reusable_template", "none"])
    route.add_argument("--director-plan")
    route.add_argument("--host", choices=["codex", "hermes", "claude-code", "generic"], default="codex")
    route.add_argument("--move", action="store_true", help="Move source files instead of copying them")
    route.set_defaults(func=cmd_route)
    return parser


def main(argv: list[str] | None = None) -> int:
    try:
        args = build_parser().parse_args(argv)
        return args.func(args)
    except (PackageError, HostError) as exc:
        print(f"error: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
