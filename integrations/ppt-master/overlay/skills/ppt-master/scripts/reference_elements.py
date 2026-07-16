#!/usr/bin/env python3
"""Extract a project-scoped design language from a reference PPTX."""

from __future__ import annotations

import argparse
import hashlib
import json
import statistics
import subprocess
import sys
from collections import Counter
from pathlib import Path
from typing import Any


class ReferenceElementsError(RuntimeError):
    """Raised when reference-element extraction cannot produce a safe lock."""


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _run_intake(pptx: Path, output_dir: Path) -> tuple[Path, Path]:
    intake = Path(__file__).resolve().parent / "pptx_intake.py"
    result = subprocess.run(
        [sys.executable, str(intake), str(pptx), "-o", str(output_dir)],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if result.returncode != 0:
        raise ReferenceElementsError((result.stderr or result.stdout).strip() or "PPTX intake failed")
    identity = output_dir / f"{pptx.stem}.identity.json"
    library = output_dir / f"{pptx.stem}.slide_library.json"
    if not identity.is_file() or not library.is_file():
        raise ReferenceElementsError("PPTX intake did not produce identity and slide-library files")
    return identity, library


def _top_values(items: list[dict[str, Any]], limit: int) -> list[Any]:
    return [item.get("value") for item in items[:limit] if item.get("value") not in {None, ""}]


def _shape_profile(pptx_path: Path) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ReferenceElementsError("python-pptx is required for template analysis") from exc

    presentation = Presentation(str(pptx_path))
    shape_types: Counter[str] = Counter()
    auto_shapes: Counter[str] = Counter()
    lines = pictures = text_shapes = 0
    per_slide: list[int] = []
    for slide in presentation.slides:
        per_slide.append(len(slide.shapes))
        for shape in slide.shapes:
            type_name = str(getattr(shape, "shape_type", "unknown"))
            shape_types[type_name] += 1
            if getattr(shape, "has_text_frame", False):
                text_shapes += 1
            if "PICTURE" in type_name:
                pictures += 1
            if "LINE" in type_name:
                lines += 1
            try:
                auto_type = getattr(shape, "auto_shape_type", None)
            except (ValueError, AttributeError):
                auto_type = None
            if auto_type is not None:
                auto_shapes[str(auto_type)] += 1
    return {
        "shape_types": dict(shape_types.most_common(12)),
        "auto_shapes": dict(auto_shapes.most_common(12)),
        "line_count": lines,
        "picture_count": pictures,
        "text_shape_count": text_shapes,
        "shapes_per_slide": {
            "median": round(statistics.median(per_slide), 1) if per_slide else 0,
            "max": max(per_slide, default=0),
        },
    }


def derive_reference_analysis(
    identity: dict[str, Any],
    library: dict[str, Any],
    *,
    source_path: Path,
    shape_profile: dict[str, Any],
) -> dict[str, Any]:
    theme = identity.get("theme") or {}
    observed = identity.get("observed") or {}
    palette = theme.get("palette") or {}
    observed_colors = _top_values(observed.get("colors") or [], 12)
    accent_colors = [
        color for color in observed_colors
        if str(color).upper() not in {"#000000", "#FFFFFF", "#F0F0F0"}
    ][:8]
    if not accent_colors:
        accent_colors = [palette.get("primary", "#3271CC")]

    fonts = theme.get("fonts") or {}
    observed_fonts = observed.get("fonts") or {}
    font_pool = list(dict.fromkeys(
        [
            (fonts.get("title") or {}).get("ea"),
            (fonts.get("title") or {}).get("latin"),
            (fonts.get("body") or {}).get("ea"),
            (fonts.get("body") or {}).get("latin"),
            *_top_values(observed_fonts.get("ea") or [], 5),
            *_top_values(observed_fonts.get("latin") or [], 5),
        ]
    ))
    font_pool = [font for font in font_pool if font]

    canvas = identity.get("canvas") or library.get("canvas_px") or {}
    width = int(canvas.get("width_px") or canvas.get("width") or 1280)
    height = int(canvas.get("height_px") or canvas.get("height") or 720)
    slides = library.get("slides") or []
    page_types = Counter(slide.get("page_type", "unknown") for slide in slides)
    title_y: list[float] = []
    body_y: list[float] = []
    slot_counts: list[int] = []
    text_lengths: list[int] = []
    for slide in slides:
        slots = slide.get("slots") or []
        slot_counts.append(len(slots))
        text_lengths.append(len(slide.get("text_summary") or ""))
        for slot in slots:
            geometry = slot.get("geometry") or {}
            y = float(geometry.get("y") or 0)
            role = slot.get("role") or ""
            if role == "title_candidate":
                title_y.append(y)
            elif role == "body_candidate":
                body_y.append(y)

    ranked = sorted(
        slides,
        key=lambda slide: (
            len(slide.get("charts") or []) + len(slide.get("diagrams") or []) + len(slide.get("tables") or []),
            len(slide.get("slots") or []),
            len(slide.get("text_summary") or ""),
        ),
        reverse=True,
    )
    representatives = [
        {
            "slide_index": slide.get("slide_index"),
            "page_type": slide.get("page_type"),
            "slot_count": len(slide.get("slots") or []),
            "relationship_objects": sum(
                len(slide.get(key) or []) for key in ("charts", "diagrams", "tables")
            ),
        }
        for slide in ranked[:8]
    ]

    theme_sizes = theme.get("sizes") or {}
    return {
        "schema_version": "1.0",
        "mode": "reference_elements",
        "source": str(source_path.resolve()),
        "source_sha256": _sha256(source_path),
        "canvas": {"width": width, "height": height, "aspect": round(width / height, 4)},
        "color_language": {
            "theme": palette,
            "observed_accents": accent_colors,
            "background_behavior": "light neutral canvas with blue-led structural accents",
        },
        "typography": {
            "font_pool": font_pool,
            "title_pt": float(theme_sizes.get("title") or 28),
            "body_pt": float(theme_sizes.get("body") or 18),
            "observed_sizes_pt": _top_values(observed.get("sizes_pt") or [], 10),
        },
        "layout_grammar": {
            "page_types": dict(page_types),
            "median_title_y": round(statistics.median(title_y), 1) if title_y else None,
            "median_body_y": round(statistics.median(body_y), 1) if body_y else None,
            "median_slot_count": round(statistics.median(slot_counts), 1) if slot_counts else 0,
            "median_text_characters": round(statistics.median(text_lengths), 1) if text_lengths else 0,
            "rhythm": "alternate anchor, structured-content, and breathing pages; do not repeat one grid",
        },
        "shape_language": shape_profile,
        "representative_pages": representatives,
        "constraints": {
            "pptx_structure_mode": "flat",
            "reuse_source_pages": False,
            "reuse_source_layouts": False,
            "inherit_page_count": False,
            "inherit_coordinates": False,
        },
    }


def _design_spec(analysis: dict[str, Any]) -> str:
    canvas = analysis["canvas"]
    colors = analysis["color_language"]
    typography = analysis["typography"]
    shapes = analysis["shape_language"]
    accents = ", ".join(f"`{value}`" for value in colors["observed_accents"])
    fonts = " / ".join(typography["font_pool"][:6])
    reps = ", ".join(str(item["slide_index"]) for item in analysis["representative_pages"])
    return f"""# Project Design System: Reference Elements

> Reference source: `{Path(analysis['source']).name}`
> Intent: extract the visual language and redesign every page; never fill or reproduce source layouts.

## Canvas

- Size: {canvas['width']} x {canvas['height']} ({canvas['aspect']}:1)
- Structure mode: `flat`
- Source page count and coordinates are evidence only, never constraints.

## Color Language

- Background: light neutral canvas, white as the primary field.
- Primary structural color: `{colors['theme'].get('primary', '#3271CC')}`.
- Observed supporting accents: {accents}.
- Use blue for hierarchy and navigation; reserve red for risk or exception semantics.
- Avoid decorative color proliferation and low-contrast pale text.

## Typography

- Safe font pool derived from the reference: {fonts}.
- Title baseline: {typography['title_pt']} pt source evidence; adapt size to the new composition.
- Body baseline: {typography['body_pt']} pt source evidence; preserve leadership-room readability.
- Use strong title/body contrast; do not inherit source text boxes or exact coordinates.

## Background And Decoration

- Prefer clean white or very light gray fields with blue structural bands, rules, or numbered anchors.
- Reuse the reference's restrained corporate-government tone, not its literal ornaments.
- Shape evidence: median {shapes['shapes_per_slide']['median']} shapes per source slide, {shapes['line_count']} lines and {shapes['picture_count']} pictures across the deck.
- Decorative elements must support hierarchy, grouping, or flow; never become a page-filling motif.

## Layout Grammar

- Build each page from its information relationship: comparison, timeline, architecture, process, matrix, ecosystem, or roadmap.
- Alternate anchor, dense, and breathing pages to preserve presentation rhythm.
- Keep one core conclusion and one dominant visual anchor per page.
- Do not use the same structure on three consecutive pages.
- Do not default to title plus three cards.

## Image Behavior

- Images are atmosphere or evidence, not a mandatory full-page background.
- Favor asymmetric image/text balance and controlled cropping.
- Preserve generous safe margins and avoid placing critical text over busy imagery.

## Evidence Boundary

- Representative source slides inspected for visual-language evidence: {reps}.
- The source deck contributes color, typography, shape, density, and rhythm only.
- Source masters, layouts, placeholders, page roster, wording, and coordinates are explicitly excluded.
"""


def _spec_lock(analysis: dict[str, Any]) -> str:
    canvas = analysis["canvas"]
    colors = analysis["color_language"]
    palette = colors["theme"]
    accents = colors["observed_accents"]
    typography = analysis["typography"]
    font_pool = typography["font_pool"] or ["Microsoft YaHei", "Arial"]
    font_stack = ", ".join(f'"{font}"' if " " in font else font for font in font_pool[:4])
    primary = palette.get("primary") or accents[0]
    secondary = palette.get("accent2") or (accents[1] if len(accents) > 1 else primary)
    title_px = max(36, round(float(typography["title_pt"]) * 4 / 3))
    body_px = max(22, round(float(typography["body_pt"]) * 4 / 3))
    return f"""# Execution Lock

## canvas
- viewBox: 0 0 {canvas['width']} {canvas['height']}
- format: PPT 16:9

## template_intent
- mode: reference_elements
- source_sha256: {analysis['source_sha256']}
- rule: inherit visual language only; redesign every page

## colors
- bg: {palette.get('background', '#FFFFFF')}
- secondary_bg: {palette.get('background_alt', '#F0F0F0')}
- primary: {primary}
- accent: {secondary}
- body: {palette.get('text', '#000000')}
- text_secondary: {palette.get('text_alt', '#768395')}
- warning: {palette.get('accent5', '#D63232')}

## typography
- font_family: {font_stack}, sans-serif
- body: {body_px}
- title: {title_px}
- subtitle: {max(body_px + 4, title_px - 8)}
- annotation: {max(16, body_px - 6)}
- footnote: {max(14, body_px - 8)}

## visual_language
- background: light neutral, blue-led structure
- decoration: restrained rules, bands, numbered anchors, selective rounded geometry
- image_behavior: asymmetric evidence or atmosphere, never mandatory wallpaper
- rhythm: alternate anchor, dense, and breathing pages

## pptx_structure
- mode: flat

## forbidden_template_reuse
- source masters
- source layouts
- source placeholders
- source page roster
- source coordinates
- direct page filling
"""


def extract_reference_elements(
    project_path: str | Path,
    pptx_path: str | Path,
    *,
    force: bool = False,
) -> dict[str, Any]:
    project = Path(project_path).expanduser().resolve()
    pptx = Path(pptx_path).expanduser().resolve()
    if not project.is_dir():
        raise ReferenceElementsError(f"project not found: {project}")
    if not pptx.is_file() or pptx.suffix.lower() != ".pptx":
        raise ReferenceElementsError(f"reference PPTX not found: {pptx}")
    design_spec = project / "design_spec.md"
    spec_lock = project / "spec_lock.md"
    if not force and (design_spec.exists() or spec_lock.exists()):
        raise ReferenceElementsError("design_spec.md or spec_lock.md already exists; pass --force to replace")

    analysis_dir = project / "analysis" / "template_reference"
    identity_path, library_path = _run_intake(pptx, analysis_dir)
    identity = json.loads(identity_path.read_text(encoding="utf-8"))
    library = json.loads(library_path.read_text(encoding="utf-8"))
    analysis = derive_reference_analysis(
        identity,
        library,
        source_path=pptx,
        shape_profile=_shape_profile(pptx),
    )
    _write_json(analysis_dir / "reference_elements.json", analysis)
    design_spec.write_text(_design_spec(analysis), encoding="utf-8")
    lock_text = _spec_lock(analysis)
    contract_path = project / "analysis" / "director_contract.json"
    if contract_path.is_file():
        contract = json.loads(contract_path.read_text(encoding="utf-8"))
        profile_id = str((contract.get("profile") or {}).get("id") or "")
        body = 20 if profile_id in {"government_strategy", "decision_meeting"} else 18
        marker = f"minimum_font_sizes: body={body}px supporting=16px footnote=12px"
        if marker not in lock_text:
            lock_text = marker + "\n\n" + lock_text
    spec_lock.write_text(lock_text, encoding="utf-8")
    return {
        "mode": "reference_elements",
        "project": str(project),
        "analysis": str(analysis_dir / "reference_elements.json"),
        "design_spec": str(design_spec),
        "spec_lock": str(spec_lock),
        "pptx_structure": "flat",
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Extract reference-PPTX visual language")
    parser.add_argument("project")
    parser.add_argument("pptx")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args(argv)
    try:
        result = extract_reference_elements(args.project, args.pptx, force=args.force)
    except ReferenceElementsError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
